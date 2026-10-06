#!/usr/bin/env python3
"""Sinh authority-matrix.js từ file Excel ma trận thẩm quyền của Tập đoàn.

Cách dùng:
    pip install openpyxl
    python3 tools/build_matrix.py "Chuan_hoa_Tu_viet_tat_RACI_TCCB_MEDVN.xlsx"

Đọc 2 sheet:
- "Authority Matrix - Toàn TĐ": luồng Đề xuất → Duyệt → Phê chuẩn theo Ban/Khối. Chỉ giữ
  các cột cần để xác định ai trong BTGĐ nhận văn bản (cột BTGĐ, TGĐ, HĐQT), hoặc cấp nào
  phê chuẩn khi văn bản không cần trình BTGĐ.
- "Từ viết tắt": mục A-D (chức danh, pháp nhân, khối/ban, nghiệp vụ) để AI hiểu chữ viết tắt.
Kết quả là chuỗi văn bản gọn, gửi kèm cho AI ở mỗi lần phân loại.
"""

import json
import os
import re
import sys

import openpyxl

MATRIX_SHEET = 'Authority Matrix - Toàn TĐ'
ABBR_SHEET = 'Từ viết tắt'
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'authority-matrix.js')

APPROVE = re.compile(r'Phê chuẩn|Ký chính')

# Cột (1-based) của sheet ma trận
C_TT, C_BAN, C_HM, C_LOAI, C_PV, C_NB, C_ND, C_TM, C_BTGD, C_TGD, C_HDQT = 1, 2, 3, 4, 6, 7, 8, 9, 10, 11, 12


def clean(v):
    if v is None:
        return ''
    return re.sub(r'\s+', ' ', str(v).replace('\n', ' / ')).strip()


def build_matrix(wb):
    ws = wb[MATRIX_SHEET]
    full_rows = {m.min_row for m in ws.merged_cells.ranges if m.min_col == 1 and m.max_col >= 12}
    lines = []
    hm = loai = ''
    rows = 0
    for r in range(6, ws.max_row + 1):
        cell = lambda c: clean(ws.cell(r, c).value)
        if r in full_rows:
            title = cell(1)
            if not title or title.startswith('KHOẢNG TRỐNG'):
                continue  # ghi chú rà soát, không dùng để phân loại
            if title.startswith('▌'):
                lines.append('')
                lines.append('## ' + title.lstrip('▌ ').strip())
            else:
                lines.append('### ' + (title if len(title) <= 160 else title[:157] + '...'))
            continue
        if not any(cell(c) for c in range(1, 13)):
            continue
        if cell(C_HM):
            hm, loai = cell(C_HM), cell(C_LOAI)
            tt = cell(C_TT)
            lines.append(f'- [{tt}] {hm}' + (f' ({loai})' if loai else '') if tt else f'- {hm}' + (f' ({loai})' if loai else ''))
        elif cell(C_LOAI) and cell(C_LOAI) != loai:
            loai = cell(C_LOAI)
        rows += 1
        pv = cell(C_PV) or 'Chung'
        btgd, tgd, hdqt = cell(C_BTGD), cell(C_TGD), cell(C_HDQT)
        # Cột BTGĐ đôi khi ghi Ban/Phòng (ví dụ "Ban NQ: Phê chuẩn") - không phải thành viên BTGĐ
        below = odd = ''
        if btgd and re.match(r'^(Ban|Phòng)\b', btgd) and not re.search(r'TGĐ', btgd):
            if APPROVE.search(btgd):
                below = btgd
            else:
                odd = btgd
            btgd = ''
        parts = []
        if btgd:
            parts.append('BTGĐ: ' + btgd)
        if tgd:
            parts.append('TGĐ: ' + tgd)
        if hdqt:
            parts.append('HĐQT: ' + hdqt)
        if parts:
            lines.append(f'  · {pv} → ' + '; '.join(parts))
        else:
            # Không có vai trò ở cấp BTGĐ/TGĐ/HĐQT: ghi lại cấp phê chuẩn cuối ở cấp dưới;
            # không thấy ai phê chuẩn thì là ma trận chưa có dữ liệu, KHÔNG được hiểu là "không trình"
            where = below
            for c, label in ((C_TM, 'Ban/Đơn vị tham mưu'), (C_ND, 'Ngành dọc quản lý'), (C_NB, 'Cấp nội bộ đơn vị')):
                v = cell(c)
                if not where and APPROVE.search(v):
                    where = v if ':' in v else f'{label}: {v}'
            if where:
                lines.append(f'  · {pv} → KHÔNG trình BTGĐ (phê chuẩn: {where})')
            else:
                lines.append(f'  · {pv} → CHƯA RÕ cấp phê chuẩn trong ma trận' + (f' (cột BTGĐ ghi: {odd})' if odd else ''))
    return '\n'.join(lines).strip(), rows


def build_abbr(wb):
    ws = wb[ABBR_SHEET]
    lines = []
    keep = False
    for row in ws.iter_rows(values_only=True):
        vals = [clean(v) for v in row]
        first = vals[0] if vals else ''
        if re.match(r'^[A-F]\. ', first):
            keep = first[0] in 'ABCD'
            if keep:
                lines.append(first)
            continue
        if not keep or not first or first == 'Từ viết tắt':
            continue
        full = vals[1] if len(vals) > 1 else ''
        if full:
            lines.append(f'{first} = {full}')
    return '\n'.join(lines)


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    src = sys.argv[1]
    wb = openpyxl.load_workbook(src, data_only=True)
    matrix, rows = build_matrix(wb)
    abbr = build_abbr(wb)
    js = (
        '// TỰ SINH bởi tools/build_matrix.py từ file ' + os.path.basename(src) + ' - không sửa tay.\n'
        '// Ma trận thẩm quyền (Authority Matrix) và từ viết tắt, gửi kèm cho AI khi phân loại.\n\n'
        f'var AUTHORITY_MATRIX_SOURCE = {json.dumps(os.path.basename(src), ensure_ascii=False)};\n'
        f'var AUTHORITY_MATRIX_ROWS = {rows};\n'
        f'var AUTHORITY_MATRIX = {json.dumps(matrix, ensure_ascii=False)};\n'
        f'var ABBREVIATIONS = {json.dumps(abbr, ensure_ascii=False)};\n'
    )
    with open(OUT, 'w', encoding='utf-8') as f:
        f.write(js)
    print(f'Đã ghi {os.path.normpath(OUT)}: {rows} dòng ma trận, {len(matrix)} ký tự; từ viết tắt {len(abbr)} ký tự')


if __name__ == '__main__':
    main()
